#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Переносит raw_data.json (результат collect.py) в лист «Данные»
файла Wordstat_Chita_analiz.xlsx. Формулы на листах «Свод» и «Динамика»
пересчитаются сами при открытии в Excel.

Запуск:
    python to_xlsx.py Wordstat_Chita_analiz.xlsx
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

from keywords import all_phrases

RAW = Path(__file__).parent / "raw_data.json"

# Колонка листа «Данные» для каждого региона (широкое соответствие)
COL = {
    "Чита": 3,
    "Забайкальский край": 6,
    "Иркутск": 8,
}


def main():
    if len(sys.argv) < 2:
        sys.exit("Укажите путь к xlsx: python to_xlsx.py Wordstat_Chita_analiz.xlsx")

    xlsx = Path(sys.argv[1])
    if not xlsx.exists():
        sys.exit(f"Файл не найден: {xlsx}")
    if not RAW.exists():
        sys.exit(f"Нет данных: {RAW}. Сначала запустите collect.py")

    data = json.loads(RAW.read_text(encoding="utf-8"))

    # Строка листа для каждой фразы
    row_of = {}
    for i, (_block, phrase) in enumerate(all_phrases()):
        row_of[phrase] = i + 2

    wb = load_workbook(xlsx)
    ws = wb["Данные"]

    written, skipped = 0, 0
    for rec in data.get("frequency", []):
        if rec.get("error") or rec.get("total_count") is None:
            skipped += 1
            continue
        r = row_of.get(rec["phrase"])
        c = COL.get(rec["region"])
        if not r or not c:
            skipped += 1
            continue
        ws.cell(row=r, column=c, value=int(rec["total_count"]))
        written += 1

    ws_i = wb["Инструкция"]
    for row in ws_i.iter_rows(min_col=2, max_col=3):
        if row[0].value == "Дата сбора:":
            row[1].value = data.get("meta", {}).get("collected_at", "")
        if row[0].value == "Кто собирал:":
            row[1].value = "Wordstat API, автоматический сбор"

    wb.save(xlsx)
    print(f"Записано значений: {written}, пропущено: {skipped}")
    print(f"Файл обновлён: {xlsx}")
    print("Динамику перенести вручную с листа «Динамика» — структура ответа API "
          "по периодам зависит от версии, проверьте raw_data.json.")


if __name__ == "__main__":
    main()
